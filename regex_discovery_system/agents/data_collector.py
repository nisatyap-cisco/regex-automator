"""Agent 1 — Data Collector (DB-only).

Sourcing:
  1. db lookup — checks db/db_index.json for a matching reference file

If the DB has no data, the pipeline reports insufficient data.
The output is normalised into:
  - raw_examples.txt   (all unique values)
  - train.txt          (100% to both Agent 2 and Agent 4)
  - test.txt

The collector also returns metadata so downstream agents know the source.
"""
from __future__ import annotations

import json
import logging
import os
import random
from pathlib import Path

logger = logging.getLogger(__name__)

SEED = 42
DB_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "db")


# ── helpers ───────────────────────────────────────────────────────────────────

def _write_lines(path: str, lines: list[str]) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w") as f:
        f.write("\n".join(lines) + "\n")


def _split_and_write(unique: list[str], source: str, paths: dict[str, str]) -> dict:
    rng = random.Random(SEED)
    rng.shuffle(unique)

    raw_path = paths.get("raw", "data/raw_examples.txt")
    train_path = paths.get("train", "data/train.txt")
    test_path = paths.get("test", "data/test.txt")

    train = list(unique)
    test = list(unique)
    logger.info("Split: %d unique → both Agent 2 & Agent 4 get all %d values", len(unique), len(unique))

    _write_lines(raw_path, unique)
    _write_lines(train_path, train)
    _write_lines(test_path, test)

    return {"total": len(unique), "train": len(train), "test": len(test)}


# ── source 1: db lookup ──────────────────────────────────────────────────────

def _match_score(condition: str, keywords: list[str]) -> int:
    """Count how many keywords appear (case-insensitive) in the condition."""
    cond_lower = condition.lower()
    return sum(1 for kw in keywords if kw.lower() in cond_lower)


def _try_db_lookup(condition: str) -> tuple[list[str] | None, str | None]:
    """Search db/db_index.json for a file whose keywords match the condition.

    Returns (values_list, description) or (None, None).
    """
    index_path = os.path.join(DB_DIR, "db_index.json")
    if not os.path.isfile(index_path):
        logger.info("Agent 1 [db]: no db_index.json found — skipping")
        return None, None

    with open(index_path) as f:
        index = json.load(f)

    best_entry = None
    best_score = 0
    for entry in index.get("entries", []):
        score = _match_score(condition, entry.get("keywords", []))
        if score > best_score:
            best_score = score
            best_entry = entry

    if best_entry is None or best_score == 0:
        logger.info("Agent 1 [db]: no matching entry for '%s'", condition)
        return None, None

    file_path = os.path.join(DB_DIR, best_entry["file"])
    if not os.path.isfile(file_path):
        logger.warning("Agent 1 [db]: file '%s' referenced but not found", file_path)
        return None, None

    fmt = best_entry.get("format", "").lower()
    col = best_entry.get("column", None)
    desc = best_entry.get("description", "")

    values: list[str] = []

    if fmt == "xlsx":
        try:
            import openpyxl
        except ImportError:
            logger.warning("Agent 1 [db]: openpyxl not installed — cannot read xlsx")
            return None, None

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col_idx = 0
        if col and col in headers:
            col_idx = headers.index(col)

        for row in ws.iter_rows(min_row=2, max_col=col_idx + 1):
            cell = row[col_idx]
            if cell.value is not None:
                values.append(str(cell.value).strip())
        wb.close()

    elif fmt in ("txt", "csv"):
        with open(file_path) as f:
            for line in f:
                v = line.strip()
                if v:
                    values.append(v)
    else:
        logger.warning("Agent 1 [db]: unsupported format '%s'", fmt)
        return None, None

    seen: set[str] = set()
    unique: list[str] = []
    for v in values:
        if v not in seen:
            seen.add(v)
            unique.append(v)

    logger.info(
        "Agent 1 [db]: matched '%s' → %s (%d values) — %s",
        condition, best_entry["file"], len(unique), desc,
    )
    return unique, desc


# ── public API ────────────────────────────────────────────────────────────────

def collect_data(condition: str, config: dict, paths: dict[str, str] | None = None) -> dict:
    paths = paths or {}

    values, db_desc = _try_db_lookup(condition)

    if not values:
        logger.warning(
            "Agent 1: no data in db for '%s' — writing empty files, "
            "pipeline will rely on policies only",
            condition,
        )
        for key in ("raw", "train", "test"):
            p = paths.get(key)
            if p:
                _write_lines(p, [])
        return {"total": 0, "train": 0, "test": 0, "source": "insufficient_data"}

    stats = _split_and_write(values, "db", paths)
    stats["source"] = "db"
    logger.info("Agent 1 done: source=db, %d unique values", stats["total"])
    return stats
